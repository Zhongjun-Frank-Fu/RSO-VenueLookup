#!/usr/bin/env python3
"""One-time script: convert UCSD校内场地预约信息总表.xlsx into src/lib/venues.json"""

import json
import re
import sys
import os

try:
    import openpyxl
except ImportError:
    print("pip3 install openpyxl")
    sys.exit(1)

XLSX = os.path.join(os.path.dirname(__file__), "..", "..", "UCSD校内场地预约信息总表.xlsx")
OUT = os.path.join(os.path.dirname(__file__), "..", "src", "lib", "venues.json")

# Building code mapping: venue name pattern -> building code
BUILDING_MAP = {
    "Ballroom": "PRICE", "Theater": "PRICE", "Forum": "PRICE",
    "Dance Studio": "PRICE",  # Price Center Dance Studio (not Wagner)
    "John Muir": "PRICE", "Bear Room": "PRICE", "Red Shoe": "PRICE",
    "Roosevelt": "PRICE", "Marshall College Room": "PRICE",
    "Warren College Room": "PRICE", "Green Table": "PRICE",
    "Governance Chambers": "PRICE", "Snake Path": "PRICE",
    "Student Leadership": "PRICE", "Sixth College Room": "PRICE",
    "Revelle College Room": "PRICE",
    "Huerta": "STCTR", "Vera Cruz": "STCTR", "Thich Nhat Hanh": "STCTR",
    "Stage Room": "STCTR", "Multipurpose Room": "STCTR",
    "SSC Conference": "SSC", "Multipurpose Room SSC": "SSC",
    "Epstein": "EPST", "LOFT": "PRICE",
    "Library Walk": "LIBWK", "Price Center Plaza": "PRICE",
    "Matthews Quad": "MATQD", "Rupertus": "RUPLN", "Town Square": "TWNSQ",
    "DIB": "DIB", "Design & Innovation": "DIB",
    "Great Hall": "GREAT", "Asante": "ASANT",
    "Atkinson": "CALIT", "Qualcomm Institute": "CALIT",
    "Conrad Prebys": "CPMC",
    "CSE Building": "EBU3B", "EBU3B": "EBU3B",
    "Jacobs Hall": "EBU1", "Qualcomm Conf": "EBU1",
    "ASML": "SME", "SME": "SME",
    "Robert Conn": "FAH", "Franklin Antonio": "FAH",
    "Mandeville": "MANDE",
    "SV-1": "SURV", "Survivance": "SURV",
    "Jeannie": "JEANN",
    "15th Floor": "HCS15",
    "Rooftop": "ROOFTOP",
    "Porton": "NUEVO", "Nuevo": "NUEVO",
    "Mosaic": "MOS",
    "Exchange": "EXCHANGE",
    "LionTree": "LION",
    "RIMAC Annex": "RIMAC", "RIMAC Conference": "RIMAC",
    "Main Gym": "RECGM", "Recreation Gym": "RECGM",
    "Triton Conference": "RIMAC", "Dugout": "RIMAC",
    "Canyonview": "CANYN",
    "Multipurpose Fields": "FIELDS",
    "Mandell Weiss Theatre": "MWEIS", "Mandell Weiss Forum": "MWEIS",
    "Potiker": "POTKR",
    "Arthur Wagner": "GH", "Wagner Theatre": "GH",
    "Shank": "SHANK",
    "Dance Studio 3": "DANCE",
}

# Booking system mapping
SYSTEM_MAP = {
    "University Centers": "university-centers",
    "T-RES": "university-centers",
    "Classroom": "classroom",
    "Registrar": "classroom",
    "HCS": "hcs",
    "Hospitality": "hcs",
    "Recreation": "recreation",
    "Rec": "recreation",
    "Theatre": "theatre-dance",
    "Dance": "theatre-dance",
}

# CSSA recommendations data
CSSA_RECS = [
    {"activity_type": "春晚 / 大型晚会", "activity_type_en": "Spring Gala / Large Events",
     "capacity_range": "500+", "advance_booking": "3-6 个月",
     "venue_names": ["Ballroom West", "Mandeville", "Ballroom East"],
     "notes": "春晚首选 AB；建议 Fall Quarter 就开始预约"},
    {"activity_type": "新生见面会 / 中型活动", "activity_type_en": "Orientation / Medium Events",
     "capacity_range": "200-500", "advance_booking": "2-3 个月",
     "venue_names": ["Ballroom East", "Ballroom West", "Great Hall"],
     "notes": "Great Hall 海景加分但需 UCSD Catering"},
    {"activity_type": "职业宣讲会", "activity_type_en": "Career Talks",
     "capacity_range": "50-200", "advance_booking": "1-2 个月",
     "venue_names": ["Forum", "Atkinson", "Jacobs"],
     "notes": "Forum 设备好；Atkinson 适合 tech talk"},
    {"activity_type": "部门例会 / 小型会议", "activity_type_en": "Team Meetings",
     "capacity_range": "10-40", "advance_booking": "1-2 周",
     "venue_names": ["Bear Room", "Red Shoe", "DIB"],
     "notes": "PC 会议室 RSO 免费；DIB 需单独预约"},
    {"activity_type": "户外集市 / Tabling", "activity_type_en": "Outdoor Fair / Tabling",
     "capacity_range": "N/A", "advance_booking": "2-4 周",
     "venue_names": ["Library Walk", "Price Center Plaza"],
     "notes": "Library Walk 人流量最大"},
    {"activity_type": "大型户外活动", "activity_type_en": "Large Outdoor Events",
     "capacity_range": "1000+", "advance_booking": "3-6 个月",
     "venue_names": ["Epstein", "Matthews Quad", "LionTree"],
     "notes": "Epstein 有专业音响；Matthews 适合集市"},
    {"activity_type": "音乐会 / 演奏会", "activity_type_en": "Concerts / Recitals",
     "capacity_range": "100-400", "advance_booking": "2-3 个月",
     "venue_names": ["Conrad Prebys", "LOFT", "Epstein"],
     "notes": "Conrad Prebys 有世界级音响"},
    {"activity_type": "体育 / 电竞活动", "activity_type_en": "Sports / Esports",
     "capacity_range": "50-5000", "advance_booking": "1-3 个月",
     "venue_names": ["Main Gym", "RIMAC Annex", "LionTree"],
     "notes": "LionTree Arena 容量最大(5000)"},
    {"activity_type": "学术讲座 / Workshop", "activity_type_en": "Academic Lectures",
     "capacity_range": "20-200", "advance_booking": "1-3 周",
     "venue_names": ["DIB", "CSE Building", "Classroom"],
     "notes": "Classroom 免费但限制多"},
    {"activity_type": "高端宴会 / VIP 活动", "activity_type_en": "VIP / Premium Events",
     "capacity_range": "50-450", "advance_booking": "2-3 个月",
     "venue_names": ["Great Hall", "SV-126", "15th Floor"],
     "notes": "均有海景；需 UCSD Catering 或 HCS"},
]


def slugify(name):
    s = re.sub(r'[^a-zA-Z0-9]+', '-', name.lower()).strip('-')
    return s[:50]


def guess_building_code(name):
    for pattern, code in BUILDING_MAP.items():
        if pattern.lower() in name.lower():
            return code
    return "UNKNOWN"


def guess_system(system_text):
    if not system_text:
        return "unknown"
    for pattern, sid in SYSTEM_MAP.items():
        if pattern.lower() in system_text.lower():
            return sid
    return "independent"


def parse_capacity(cap_text):
    if not cap_text or str(cap_text).strip() == '':
        return {}, None
    text = str(cap_text)
    cap = {}
    max_val = None
    patterns = [
        (r'Theater\s*[:\s]*(\d[\d,]*)', 'theater'),
        (r'Banquet\s*[:\s]*(\d[\d,]*)', 'banquet'),
        (r'Classroom\s*[:\s]*(\d[\d,]*)', 'classroom'),
        (r'Reception\s*[:\s]*(\d[\d,]*)', 'reception'),
        (r'Boardroom\s*[:\s]*(\d[\d,]*)', 'boardroom'),
        (r'Circle\s*[:\s]*(\d[\d,]*)', 'circle'),
        (r'U-?shape\s*[:\s]*(\d[\d,]*)', 'u_shape'),
        (r'Hollow\s*Sq\w*\s*[:\s]*(\d[\d,]*)', 'hollow_square'),
    ]
    for pat, key in patterns:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            val = int(m.group(1).replace(',', ''))
            cap[key] = val
            max_val = max(max_val or 0, val)

    if not cap:
        nums = re.findall(r'(\d[\d,]+)', text)
        if nums:
            vals = [int(n.replace(',', '')) for n in nums]
            max_val = max(vals)
            if len(vals) == 1:
                cap['max'] = vals[0]

    return cap, max_val


def parse_rates(rso_text, dept_text, ext_text):
    def clean(t):
        return str(t).strip() if t and str(t).strip() not in ('nan', 'None', '') else None
    rso = clean(rso_text)
    dept = clean(dept_text)
    ext = clean(ext_text)
    return {
        "rso": rso or "需联系咨询",
        "department": dept or "需联系咨询",
        "external": ext or "需联系咨询"
    }


def is_rso_free(rso_text):
    if not rso_text:
        return False
    text = str(rso_text)
    return '免费' in text or 'free' in text.lower()


def extract_equipment(equip_text):
    if not equip_text or str(equip_text).strip() in ('nan', 'None', ''):
        return []
    items = re.split(r'[;；、\n]+', str(equip_text))
    return [i.strip() for i in items if i.strip()]


def clean_str(val):
    if val is None or str(val).strip() in ('nan', 'None', ''):
        return None
    return str(val).strip()


def clean_name(val):
    """Clean venue name: collapse newlines to spaces."""
    s = clean_str(val)
    if s:
        s = re.sub(r'\s*\n\s*', ' ', s)
    return s


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # Sheet 1: Booking systems overview
    ws1 = wb['预约系统总览']
    booking_systems = []
    system_ids = ["university-centers", "classroom", "independent", "hcs", "recreation", "theatre-dance"]
    system_colors = {
        "university-centers": "#2563EB",
        "classroom": "#6B7280",
        "independent": "#8B0000",
        "hcs": "#7C3AED",
        "recreation": "#059669",
        "theatre-dance": "#EA580C"
    }
    for i, row in enumerate(ws1.iter_rows(min_row=2, values_only=True)):
        if i >= 6:
            break
        name = clean_str(row[0])
        if not name:
            continue
        sid = system_ids[i] if i < len(system_ids) else f"system-{i}"
        short = name.split('\n')[0].strip() if name else ""
        booking_systems.append({
            "id": sid,
            "name": name.replace('\n', ' '),
            "short_name": short[:30],
            "color": system_colors.get(sid, "#666"),
            "scope": clean_str(row[1]) or "",
            "method": clean_str(row[2]) or "",
            "contact": clean_str(row[3]) or "",
            "timeline": clean_str(row[4]) or "",
            "applicable_to": clean_str(row[5]) or "",
            "rate_overview": clean_str(row[6]) or "",
            "notes": clean_str(row[7]) or ""
        })

    # Sheet 2: Venue details
    ws2 = wb['场地详细信息']
    headers = [cell.value for cell in next(ws2.iter_rows(min_row=1, max_row=1))]

    venues = {}  # building_code -> [spaces]
    current_category = None
    category_order = []
    venue_id_map = {}  # for CSSA recommendation matching

    for row in ws2.iter_rows(min_row=2, values_only=True):
        name = clean_name(row[0])
        if not name:
            continue

        # Category header row
        if name.startswith('▶'):
            current_category = name.replace('▶', '').strip()
            category_order.append(current_category)
            continue

        building_code = guess_building_code(name)
        system_id = guess_system(clean_str(row[2]))
        cap, max_cap = parse_capacity(clean_str(row[5]))
        rates = parse_rates(clean_str(row[7]), clean_str(row[8]), clean_str(row[9]))
        rso_free = is_rso_free(clean_str(row[7]))
        equipment = extract_equipment(clean_str(row[11]))
        venue_id = slugify(name)

        space = {
            "id": venue_id,
            "name": name,
            "category": current_category,
            "description": clean_str(row[1]) or "",
            "system": system_id,
            "location": clean_str(row[3]) or "",
            "area_sqft": None,
            "capacity": cap,
            "max_capacity": max_cap,
            "type": clean_str(row[6]) or "",
            "rates": rates,
            "rso_free": rso_free,
            "managed_by": clean_str(row[10]) or "",
            "equipment": equipment,
            "booking_method": clean_str(row[12]) or "",
            "contact": clean_str(row[13]) or "",
            "notes": clean_str(row[14]) or "",
            "tags": [],
            "cssa_recommended": False
        }

        # Parse area
        area_raw = clean_str(row[4])
        if area_raw:
            area_num = re.sub(r'[^\d]', '', area_raw)
            if area_num:
                space["area_sqft"] = int(area_num)

        if building_code not in venues:
            venues[building_code] = []
        venues[building_code].append(space)
        venue_id_map[venue_id] = space

    # Mark CSSA recommended venues and add tags
    recs_output = []
    for rec in CSSA_RECS:
        matched_ids = []
        for vid, space in venue_id_map.items():
            for rname in rec["venue_names"]:
                if rname.lower() in space["name"].lower():
                    space["cssa_recommended"] = True
                    if rec["activity_type"].split('/')[0].strip() not in space["tags"]:
                        space["tags"].append(rec["activity_type"].split('/')[0].strip())
                    matched_ids.append(vid)
                    break
        recs_output.append({
            "activity_type": rec["activity_type"],
            "activity_type_en": rec["activity_type_en"],
            "capacity_range": rec["capacity_range"],
            "venue_ids": matched_ids,
            "advance_booking": rec["advance_booking"],
            "notes": rec["notes"]
        })

    # Build output
    output = {
        "venues": {},
        "booking_systems": booking_systems,
        "cssa_recommendations": recs_output,
        "category_order": category_order,
        "meta": {
            "version": "1.0",
            "last_updated": "2026-04",
            "total_venues": sum(len(v) for v in venues.values()),
            "total_buildings": len(venues),
            "source": "UCSD校内场地预约信息总表.xlsx"
        }
    }

    for bcode, spaces in venues.items():
        output["venues"][bcode] = {"spaces": spaces}

    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    total = sum(len(v) for v in venues.values())
    print(f"Generated {OUT}")
    print(f"  {total} venues across {len(venues)} buildings")
    print(f"  {len(booking_systems)} booking systems")
    print(f"  {len(recs_output)} CSSA recommendations")
    print(f"  Buildings: {sorted(venues.keys())}")


if __name__ == '__main__':
    main()
